from __future__ import annotations

import base64
import json
import logging
import mimetypes
import os
import re
import shutil
import urllib.request
from datetime import date
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException, UploadFile
from psycopg2.extras import RealDictCursor

from database import db_connection
from service_config import settings

logger = logging.getLogger(__name__)


TEXT_MIME_PREFIXES = ("text/",)
TEXT_MIME_TYPES = {
    "application/json",
    "application/xml",
    "application/csv",
    "application/vnd.ms-excel",
}
IMAGE_MIME_PREFIXES = ("image/",)
MEMORY_REFRESH_EVERY_MESSAGES = 6
MAX_MEMORY_CHARS = 4500
MEMORY_LLM_COMPRESSION_TRIGGER = 2600
MEMORY_LLM_TARGET_CHARS = 1600

GIA_INPUT_TOKEN_COST_EUR_PER_1K = 0.00014
GIA_OUTPUT_TOKEN_COST_EUR_PER_1K = 0.00056
GIA_IMAGE_GENERATION_COST_EUR = 0.04


class GiaService:
    @staticmethod
    def create_conversation(user_id: str, title: str | None = None) -> dict:
        with db_connection() as connection:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                GiaService._ensure_schema(cursor)
                cursor.execute(
                    """
                    INSERT INTO gia_conversaciones (user_id, titulo)
                    VALUES (%s, COALESCE(NULLIF(%s, ''), 'Nueva conversación'))
                    RETURNING id::text, titulo, memoria_resumen, created_at, updated_at
                    """,
                    (user_id, title),
                )
                row = dict(cursor.fetchone())
                connection.commit()
        return row

    @staticmethod
    def list_conversations(user_id: str) -> list[dict]:
        with db_connection() as connection:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                GiaService._ensure_schema(cursor)
                cursor.execute(
                    """
                    SELECT
                        c.id::text,
                        c.titulo,
                        c.updated_at,
                        COUNT(m.id) AS message_count,
                        (
                            SELECT m2.content
                            FROM gia_mensajes m2
                            WHERE m2.conversacion_id = c.id
                            ORDER BY m2.created_at DESC
                            LIMIT 1
                        ) AS last_message
                    FROM gia_conversaciones c
                    LEFT JOIN gia_mensajes m ON m.conversacion_id = c.id
                    WHERE c.user_id = %s AND c.archivada = FALSE
                    GROUP BY c.id
                    ORDER BY c.updated_at DESC
                    """,
                    (user_id,),
                )
                rows = [GiaService._conversation_summary(row) for row in cursor.fetchall()]
        return rows

    @staticmethod
    def get_conversation(user_id: str, conversation_id: str) -> dict | None:
        with db_connection() as connection:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                GiaService._ensure_schema(cursor)
                conversation = GiaService._get_owned_conversation(cursor, user_id, conversation_id)
                if not conversation:
                    return None
                messages = GiaService._get_messages(cursor, conversation_id)
                files = GiaService._get_files(cursor, conversation_id)
        return {
            **conversation,
            "messages": messages,
            "files": files,
        }

    @staticmethod
    def send_message(
        user_id: str,
        conversation_id: str,
        message: str,
        mode: str = "respuesta",
        files: list[UploadFile] | None = None,
    ) -> dict:
        if mode not in {"respuesta", "pdf", "imagen"}:
            raise HTTPException(status_code=422, detail="Modo GIA no soportado")

        uploaded = files or []
        with db_connection() as connection:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                GiaService._ensure_schema(cursor)
                conversation = GiaService._get_owned_conversation(cursor, user_id, conversation_id)
                if not conversation:
                    raise HTTPException(status_code=404, detail="Conversación no encontrada")

                user_message = GiaService._insert_message(cursor, conversation_id, "user", message)
                attachments = [
                    GiaService._store_upload(cursor, conversation_id, user_message["id"], user_id, upload)
                    for upload in uploaded
                    if upload.filename
                ]

                GiaService._touch_conversation(cursor, conversation_id, message)
                history = GiaService._get_messages(cursor, conversation_id, limit=40)
                conversation_files = GiaService._get_files_for_context(cursor, conversation_id)
                connection.commit()

            effective_mode = GiaService._effective_generation_mode(message, mode, attachments)

            assistant_text, generated_image, usage, suggested_filename = GiaService._run_openai(
                message,
                history,
                attachments,
                effective_mode,
                conversation_files,
                conversation.get("memoria_resumen") or "",
            )

        generated_files: list[dict] = []
        assistant_text_to_store = assistant_text
        with db_connection() as connection:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                
                if effective_mode == "pdf" and not GiaService._is_pdf_failure_answer(assistant_text):
                    generated_files.append(
                        GiaService._create_pdf(cursor, conversation_id, None, user_id, assistant_text, suggested_filename)
                    )
                    assistant_text_to_store = "Documento PDF generado correctamente." if generated_files else "No se pudo generar el PDF." 
                
                elif effective_mode == "imagen" and generated_image:
                    generated_files.append(
                        GiaService._store_generated_image(
                            cursor,
                            conversation_id,
                            None,
                            user_id,
                            generated_image,
                            suggested_filename,
                        )
                    )
                    assistant_text_to_store = "Imagen generada correctamente." if generated_files else "No se pudo generar la imagen."
                
                elif effective_mode == "pdf" and GiaService._is_pdf_failure_answer(assistant_text):
                    assistant_text_to_store = assistant_text
                
                elif effective_mode == "imagen" and not generated_image:
                    assistant_text_to_store = assistant_text
                
                else:
                    assistant_text_to_store = assistant_text

                
                assistant_message = GiaService._insert_message(
                    cursor,
                    conversation_id,
                    "assistant",
                    assistant_text_to_store,
                    {
                        "mode": effective_mode,
                        "requested_mode": mode,
                        "usage": usage,
                    },
                )
                GiaService._record_usage(
                    cursor,
                    user_id=user_id,
                    conversation_id=conversation_id,
                    message_id=assistant_message["id"],
                    mode=effective_mode,
                    requested_mode=mode,
                    usage=usage,
                )
                
                for f in generated_files:
                    cursor.execute(
                        """
                        UPDATE gia_archivos SET mensaje_id = %s WHERE id = %s
                        """,
                        (assistant_message["id"], f["id"])
                    )
                GiaService._touch_conversation(cursor, conversation_id, message)
                summary = GiaService._get_conversation_summary(cursor, conversation_id)
                if GiaService._should_refresh_memory(summary.get("message_count"), conversation.get("memoria_resumen")):
                    GiaService._refresh_conversation_memory(cursor, conversation_id)
                user_message["files"] = GiaService._files_for_message(cursor, user_message["id"])
                assistant_message["files"] = GiaService._files_for_message(cursor, assistant_message["id"])
                connection.commit()

        return {
            "conversation": summary,
            "user_message": user_message,
            "assistant_message": assistant_message,
            "generated_files": generated_files,
        }

    @staticmethod
    def _effective_generation_mode(message: str, requested_mode: str, attachments: list[dict]) -> str:
        if requested_mode != "pdf":
            return requested_mode
        if GiaService._is_pdf_generation_request(message, bool(attachments)):
            return "pdf"
        return "respuesta"

    @staticmethod
    def _is_pdf_generation_request(message: str, has_attachments: bool = False) -> bool:
        text = " ".join((message or "").lower().split())
        if not text:
            return False

        explicit_pdf_terms = (
            "pdf",
            "documento",
            "informe",
            "reporte",
            "entregable",
            "memoria",
            "dossier",
        )
        generation_verbs = (
            "genera",
            "generame",
            "genérame",
            "generar",
            "crea",
            "crear",
            "redacta",
            "redactar",
            "exporta",
            "exportar",
            "convierte",
            "convertir",
            "prepara",
            "preparar",
            "elabora",
            "elaborar",
            "haz",
            "hacer",
            "dame",
            "necesito",
        )
        reference_terms = (
            "ese informe",
            "este informe",
            "el informe",
            "lo anterior",
            "esa información",
            "esta información",
            "esa informacion",
            "esta informacion",
            "a partir de eso",
            "a partir de ese",
            "a partir de esta",
        )
        meta_questions = (
            "sobre que",
            "sobre qué",
            "qué te he preguntado",
            "que te he preguntado",
            "por que",
            "por qué",
            "porque me has dicho",
            "por qué me has dicho",
            "si lo has generado",
        )

        if any(term in text for term in meta_questions) and not any(term in text for term in explicit_pdf_terms):
            return False
        if any(term in text for term in reference_terms) and any(term in text for term in generation_verbs + ("pdf",)):
            return True
        if any(term in text for term in explicit_pdf_terms) and any(term in text for term in generation_verbs):
            return True
        if has_attachments and any(term in text for term in explicit_pdf_terms + generation_verbs):
            return True
        return len(text) >= 80 and any(term in text for term in ("analisis", "análisis", "resumen", "detalle", "detallado"))

    @staticmethod
    def _is_pdf_failure_answer(answer: str) -> bool:
        normalized = " ".join((answer or "").lower().split())
        return (
            "no dispongo de datos suficientes" in normalized
            or "no hay datos suficientes" in normalized
            or "no_puede_generar" in normalized
            or "no_hay_datos_suficientes" in normalized
        )

    @staticmethod
    def _history_without_current_user_message(history: list[dict], current_message: str) -> list[dict]:
        if not history:
            return []
        last = history[-1]
        if last.get("role") == "user" and (last.get("content") or "") == current_message:
            return history[:-1]
        return history

    @staticmethod
    def _pdf_source_context(history: list[dict]) -> str:
        source_blocks: list[str] = []
        for item in reversed(history):
            if item.get("role") != "assistant":
                continue
            content = (item.get("content") or "").strip()
            if not content or GiaService._is_pdf_failure_answer(content):
                continue
            source_blocks.append(content[:2500])
            if len(source_blocks) >= 3:
                break
        if not source_blocks:
            return ""
        source_text = "\n\n---\n\n".join(reversed(source_blocks))
        return (
            "Fuente conversacional prioritaria para generar el PDF. "
            "Usa estos datos antes de pedir más información y no los contradigas:\n\n"
            f"{source_text}"
        )[:8000]

    @staticmethod
    def delete_conversation(user_id: str, conversation_id: str) -> bool:
        with db_connection() as connection:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                GiaService._ensure_schema(cursor)
                cursor.execute(
                    """
                    DELETE FROM gia_conversaciones
                    WHERE id = %s AND user_id = %s
                    RETURNING id::text
                    """,
                    (conversation_id, user_id),
                )
                deleted = cursor.fetchone()
                connection.commit()
        if not deleted:
            return False
        try:
            storage_dir = Path(settings.gia_storage_dir) / re.sub(r"[^a-zA-Z0-9-]", "", conversation_id)
            if storage_dir.exists():
                shutil.rmtree(storage_dir, ignore_errors=True)
        except Exception:
            pass
        return True

    @staticmethod
    def get_file_for_download(user_id: str, file_id: str) -> dict | None:
        with db_connection() as connection:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                GiaService._ensure_schema(cursor)
                cursor.execute(
                    """
                    SELECT
                        a.id::text,
                        a.nombre_original,
                        a.mime_type,
                        a.ruta_archivo
                    FROM gia_archivos a
                    JOIN gia_conversaciones c ON c.id = a.conversacion_id
                    WHERE a.id = %s AND c.user_id = %s
                    """,
                    (file_id, user_id),
                )
                row = cursor.fetchone()
        return dict(row) if row else None

    @staticmethod
    def _run_openai(
        message: str,
        history: list[dict],
        attachments: list[dict],
        mode: str,
        conversation_files: list[dict] | None = None,
        conversation_memory: str = "",
    ) -> tuple[str, bytes | None, dict, str | None]:
        if not settings.openai_api_key:
            raise HTTPException(status_code=503, detail="OPENAI_API_KEY no está configurada")

        try:
            from openai import OpenAI
        except ImportError as exc:
            raise HTTPException(status_code=503, detail="Dependencia OpenAI no instalada") from exc

        client = OpenAI(api_key=settings.openai_api_key)
        context = GiaService._attachment_context(attachments, conversation_files or [])
        prompt_history = GiaService._history_without_current_user_message(history, message)
        pdf_source_context = GiaService._pdf_source_context(prompt_history) if mode == "pdf" else ""
        today = date.today().isoformat()
        usage = {
            "model": settings.openai_gia_model,
            "prompt_tokens": 0,
            "completion_tokens": 0,
            "total_tokens": 0,
            "chat_calls": 0,
            "image_generations": 0,
            "estimated_cost_eur": 0.0,
        }
        
        
        if mode == "pdf":
            mode_instruction = (
                "Redacta un documento profesional completo y estructurado listo para convertir a PDF. "
                "El documento debe tener: introducción, secciones claramente tituladas, datos y análisis detallado, "
                "y conclusiones. Usa párrafos separados para cada concepto. "
                "Los mensajes previos de esta conversación, incluidas respuestas anteriores del asistente y "
                "resultados de herramientas, son datos fuente válidos. Si el usuario pide un PDF de 'ese informe', "
                "'lo anterior' o una referencia similar, usa el último informe o análisis relevante ya presente "
                "en la conversación. "
                "IMPORTANTE: Si no encuentras datos suficientes o concretos para generar un documento real, "
                "responde ÚNICAMENTE: 'NO_HAY_DATOS_SUFICIENTES'. No intentes generar nada si faltan datos. "
                "Al final del documento, sugiere un nombre de archivo profesional y descriptivo para el PDF usando la convención [[FILENAME: nombre.pdf]]."
            )
        elif mode == "imagen":
            mode_instruction = (
                "Primero analiza si hay suficientes datos contextuales para generar una imagen significativa. "
                "Si SÍ hay datos (clientes, trabajos, métricas, etc.), proporciona una descripción detallada de la imagen. "
                "IMPORTANTE: Si NO hay datos suficientes o la solicitud es vaga, responde ÚNICAMENTE: 'NO_PUEDE_GENERAR'. "
                "Nunca inventes datos. Las imágenes deben basarse únicamente en información real del contexto. "
                "Si generas una imagen, sugiere un nombre de archivo profesional y descriptivo para la imagen usando la convención [[FILENAME: nombre.png]]."
            )
        else:
            mode_instruction = "Responde normalmente en el chat, ayudando al usuario."

        messages: list[dict] = [
            {
                "role": "system",
                "content": (
                    "Eres GIA, el portal de IA de GestorIA. Ayudas a una gestoría española con "
                    "consultas internas, análisis de archivos, redacción documental fiscal/laboral "
                    "y generación de entregables. Responde siempre en español, no inventes datos "
                    f"y usa la fecha actual {today}. {mode_instruction}"
                ),
            }
        ]
        if conversation_memory.strip():
            messages.append(
                {
                    "role": "system",
                    "content": (
                        "Memoria persistente de la conversación (hechos, decisiones y contexto acumulado):\n"
                        + conversation_memory[:MAX_MEMORY_CHARS]
                    ),
                }
            )
        if pdf_source_context:
            messages.append(
                {
                    "role": "system",
                    "content": pdf_source_context,
                }
            )
        if len(prompt_history) > 20:
            older_history = prompt_history[:-20][-10:]
            memory_lines = [
                f"- {it['role']}: {(it.get('content') or '').strip().replace(chr(10), ' ')[:180]}"
                for it in older_history
                if (it.get("content") or "").strip()
            ]
            if memory_lines:
                messages.append(
                    {
                        "role": "system",
                        "content": (
                            "Memoria resumida de esta conversación (turnos anteriores):\n"
                            + "\n".join(memory_lines)
                        ),
                    }
                )

        for item in prompt_history[-20:]:
            if item["role"] in {"user", "assistant"}:
                messages.append({"role": item["role"], "content": item["content"]})

        user_content: str | list[dict] = message + context
        image_parts = GiaService._image_parts(conversation_files or attachments)
        if image_parts:
            user_content = [{"type": "text", "text": message + context}, *image_parts]
        messages.append({"role": "user", "content": user_content})

        try:
            for _ in range(5):
                response = client.chat.completions.create(
                    model=settings.openai_gia_model,
                    messages=messages,
                    tools=__import__("routes.ai", fromlist=["TOOLS"]).TOOLS,
                    tool_choice="auto",
                    max_tokens=1800,
                    temperature=0.25,
                )
                GiaService._accumulate_chat_usage(usage, response)
                choice = response.choices[0]
                if choice.finish_reason != "tool_calls":
                    answer = choice.message.content or ""
                    break
                messages.append(choice.message)
                ai_module = __import__("routes.ai", fromlist=["_execute_tool"])
                for tool_call in choice.message.tool_calls:
                    args = json.loads(tool_call.function.arguments or "{}")
                    result = ai_module._execute_tool(tool_call.function.name, args)
                    messages.append({"role": "tool", "tool_call_id": tool_call.id, "content": result})
            else:
                response = client.chat.completions.create(
                    model=settings.openai_gia_model,
                    messages=messages,
                    max_tokens=1800,
                    temperature=0.25,
                )
                GiaService._accumulate_chat_usage(usage, response)
                answer = response.choices[0].message.content or ""

            
            is_no_data_for_pdf = mode == "pdf" and "NO_HAY_DATOS_SUFICIENTES" in answer.upper()
            is_no_data_for_image = mode == "imagen" and "NO_PUEDE_GENERAR" in answer.upper()
            
            generated_image = None
            if mode == "imagen" and not is_no_data_for_image:
                
                image_prompt = GiaService._generate_image_prompt(client, message, context, answer)
                if image_prompt and image_prompt != "NO_PUEDE_GENERAR":
                    generated_image = GiaService._generate_image(client, image_prompt)
                    if generated_image:
                        usage["image_generations"] = int(usage.get("image_generations") or 0) + 1
                    if not answer.strip():
                        answer = "He generado la imagen solicitada basada en los datos disponibles y la he adjuntado a esta conversación."
                else:
                    is_no_data_for_image = True
            
            if is_no_data_for_image:
                
                answer = "No tengo suficientes datos contextuales para generar una imagen significativa. Por favor, proporciona más información específica sobre qué deseas visualizar."
            
            
            if is_no_data_for_pdf:
                answer = "No dispongo de datos suficientes para generar un documento profesional. Por favor, proporciona más información o consulta datos específicos primero."
            
            usage["estimated_cost_eur"] = GiaService._estimate_usage_cost_eur(usage)
            
            suggested_filename = None
            match = re.search(r"\[\[FILENAME: ([^\]]+)\]\]", answer)
            if match:
                suggested_filename = match.group(1).strip()
                answer = re.sub(r"\[\[FILENAME: [^\]]+\]\]", "", answer).strip()
            return answer, generated_image, usage, suggested_filename
        except HTTPException:
            raise
        except Exception as exc:
            logger.exception("Fallo inesperado al contactar con OpenAI (modo=%s)", mode)
            raise HTTPException(
                status_code=502,
                detail=f"Error al contactar con OpenAI: {exc}",
            ) from exc

    @staticmethod
    def _accumulate_chat_usage(usage: dict, response) -> None:
        payload = getattr(response, "usage", None)
        if not payload:
            return
        prompt_tokens = int(getattr(payload, "prompt_tokens", 0) or 0)
        completion_tokens = int(getattr(payload, "completion_tokens", 0) or 0)
        total_tokens = int(getattr(payload, "total_tokens", 0) or 0)
        usage["prompt_tokens"] = int(usage.get("prompt_tokens") or 0) + prompt_tokens
        usage["completion_tokens"] = int(usage.get("completion_tokens") or 0) + completion_tokens
        usage["total_tokens"] = int(usage.get("total_tokens") or 0) + total_tokens
        usage["chat_calls"] = int(usage.get("chat_calls") or 0) + 1

    @staticmethod
    def _estimate_usage_cost_eur(usage: dict) -> float:
        prompt_tokens = float(usage.get("prompt_tokens") or 0)
        completion_tokens = float(usage.get("completion_tokens") or 0)
        image_generations = float(usage.get("image_generations") or 0)
        cost = (
            (prompt_tokens / 1000.0) * GIA_INPUT_TOKEN_COST_EUR_PER_1K
            + (completion_tokens / 1000.0) * GIA_OUTPUT_TOKEN_COST_EUR_PER_1K
            + image_generations * GIA_IMAGE_GENERATION_COST_EUR
        )
        return round(cost, 6)

    @staticmethod
    def _generate_image(client, prompt: str) -> bytes | None:
        model = settings.openai_image_model
        if model.startswith("gpt-4") or model.startswith("o1") or model.startswith("o3"):
            raise HTTPException(
                status_code=503,
                detail=(
                    f"OPENAI_IMAGE_MODEL='{model}' no es un modelo de generación de "
                    "imagen. Configura 'gpt-image-1' o 'dall-e-3'."
                ),
            )

        
        if "NO_PUEDE_GENERAR" in prompt.upper() or "no tengo suficientes datos" in prompt.lower():
            return None

        try:
            response = client.images.generate(
                model=model,
                prompt=prompt[:4000],
                size="1024x1024",
                n=1,
            )
        except Exception as exc:
            logger.exception("Fallo en client.images.generate (modelo=%s)", model)
            
            logger.warning("No se pudo generar imagen, continuando sin ella")
            return None

        data = getattr(response, "data", None) or []
        if not data:
            logger.warning("OpenAI no devolvió ninguna imagen")
            return None

        item = data[0]
        b64 = getattr(item, "b64_json", None)
        if b64:
            try:
                return base64.b64decode(b64)
            except Exception as exc:
                logger.exception("b64_json inválido en respuesta de OpenAI")
                return None

        url = getattr(item, "url", None)
        if url:
            try:
                with urllib.request.urlopen(url, timeout=30) as resp:
                    return resp.read()
            except Exception as exc:
                logger.exception("No se pudo descargar la imagen desde %s", url)
                return None

        logger.warning("OpenAI no devolvió una imagen utilizable")
        return None

    @staticmethod
    def _generate_image_prompt(client, user_message: str, context: str, initial_response: str) -> str | None:
        """
        Genera un prompt mejorado para DALL-E basado en la respuesta inicial y contexto.
        Solo genera si hay datos suficientes indicados en la respuesta.
        """
        
        if "NO_PUEDE_GENERAR" in initial_response.upper():
            return "NO_PUEDE_GENERAR"
        
        
        try:
            refine_messages = [
                {
                    "role": "system",
                    "content": (
                        "Eres un especialista en prompts para generación de imágenes. "
                        "Tu tarea es analizar la respuesta anterior y generar un prompt detallado, "
                        "visual y profesional para DALL-E 3 que transmita la información clave de manera gráfica. "
                        "El prompt debe basarse ÚNICAMENTE en datos mencionados en la respuesta y el contexto, "
                        "sin inventar nada. Si no hay datos suficientes, responde: NO_PUEDE_GENERAR"
                    )
                },
                {
                    "role": "user",
                    "content": (
                        f"Solicitud del usuario: {user_message}\n\n"
                        f"Contexto disponible: {context[:2000]}\n\n"
                        f"Respuesta generada:\n{initial_response}\n\n"
                        "Genera un prompt detallado para crear una imagen que ilustre esta información. "
                        "El prompt debe ser visual, profesional y específico (máximo 400 caracteres). "
                        "Si no hay datos suficientes para hacer una imagen meaningful, responde: NO_PUEDE_GENERAR"
                    )
                }
            ]
            
            response = client.chat.completions.create(
                model=settings.openai_gia_model,
                messages=refine_messages,
                max_tokens=500,
                temperature=0.7,
            )
            
            image_prompt = response.choices[0].message.content or ""
            return image_prompt.strip() if image_prompt.strip() else None
        except Exception as exc:
            logger.exception("Error generando prompt para imagen: %s", exc)
            return None

    @staticmethod
    def _store_upload(cursor, conversation_id: str, message_id: str, user_id: str, upload: UploadFile) -> dict:
        storage_dir = GiaService._conversation_dir(conversation_id)
        suffix = Path(upload.filename or "archivo").suffix
        stored_name = f"{uuid4().hex}{suffix}"
        path = storage_dir / stored_name
        with path.open("wb") as output:
            shutil.copyfileobj(upload.file, output)
        mime_type = upload.content_type or mimetypes.guess_type(upload.filename or "")[0] or "application/octet-stream"
        size = path.stat().st_size
        extracted = GiaService._extract_text(path, mime_type)
        file_item = GiaService._insert_file(
            cursor,
            conversation_id,
            message_id,
            user_id,
            upload.filename or stored_name,
            stored_name,
            mime_type,
            size,
            str(path),
            "upload",
            extracted,
        )
        file_item["ruta_archivo"] = str(path)
        file_item["extracted_text"] = extracted
        return file_item

    @staticmethod
    def _store_generated_image(cursor, conversation_id: str, message_id: str, user_id: str, image_bytes: bytes, suggested_filename: str | None = None) -> dict:
        storage_dir = GiaService._conversation_dir(conversation_id)
        
        if suggested_filename and re.match(r"^[\w\- .áéíóúÁÉÍÓÚñÑ]+\.(png|jpg|jpeg|webp)$", suggested_filename, re.IGNORECASE):
            original_name = suggested_filename
            suffix = Path(suggested_filename).suffix or ".png"
        else:
            original_name = "imagen-generada-gia.png"
            suffix = ".png"
        stored_name = f"gia-imagen-{uuid4().hex}{suffix}"
        path = storage_dir / stored_name
        path.write_bytes(image_bytes)
        return GiaService._insert_file(
            cursor,
            conversation_id,
            message_id,
            user_id,
            original_name,
            stored_name,
            "image/png",
            len(image_bytes),
            str(path),
            "image",
            None,
        )

    @staticmethod
    def _create_pdf(cursor, conversation_id: str, message_id: str, user_id: str, content: str, suggested_filename: str | None = None) -> dict:
        try:
            from reportlab.lib.enums import TA_LEFT, TA_CENTER
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import cm
            from reportlab.platypus import (
                HRFlowable,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
            )
        except ImportError as exc:
            logger.exception("reportlab no instalado")
            raise HTTPException(
                status_code=503,
                detail="Falta la dependencia 'reportlab' para generar PDFs.",
            ) from exc

        storage_dir = GiaService._conversation_dir(conversation_id)
        
        if suggested_filename and re.match(r"^[\w\- .áéíóúÁÉÍÓÚñÑ]+\.pdf$", suggested_filename, re.IGNORECASE):
            original_name = suggested_filename
        else:
            original_name = "documento-generado-gia.pdf"
        stored_name = f"gia-documento-{uuid4().hex}.pdf"
        path = storage_dir / stored_name

        styles = getSampleStyleSheet()
        
        
        title_style = ParagraphStyle(
            "GiaTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            spaceAfter=6,
            textColor="#1a3528",
            alignment=TA_CENTER,
        )
        subtitle_style = ParagraphStyle(
            "GiaSubtitle",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=16,
            spaceAfter=10,
            textColor="#2d5a42",
            spaceBefore=12,
        )
        meta_style = ParagraphStyle(
            "GiaMeta",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=8,
            leading=11,
            textColor="#7b8896",
            spaceAfter=16,
            alignment=TA_CENTER,
        )
        body_style = ParagraphStyle(
            "GiaBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            alignment=TA_LEFT,
            spaceAfter=10,
        )

        
        story = []
        
        
        story.append(Paragraph("GESTORIA", title_style))
        story.append(Paragraph(f"Documento profesional generado por GIA · {date.today().strftime('%d/%m/%Y')}", meta_style))
        story.append(HRFlowable(width="100%", thickness=1.2, color="#1a3528", spaceAfter=12))

        
        text = (content or "").strip() or "(Sin contenido)"
        
        
        paragraphs = re.split(r"\n\s*\n", text)
        for para_text in paragraphs:
            para_text = para_text.strip()
            if not para_text:
                continue
                
            
            if len(para_text) < 80 and not para_text.endswith(('.', ',')):
                escaped = GiaService._html_escape(para_text)
                story.append(Paragraph(escaped, subtitle_style))
            else:
                
                escaped = GiaService._html_escape(para_text).replace("\n", "<br/>")
                story.append(Paragraph(escaped, body_style))
                story.append(Spacer(1, 4))

        
        story.append(Spacer(1, 12))
        story.append(HRFlowable(width="100%", thickness=0.6, color="#dce8e0", spaceAfter=6))
        story.append(Paragraph(
            "Documento confidencial generado por GestorIA · GIA. Uso exclusivo interno.",
            meta_style
        ))

        doc = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            leftMargin=2 * cm,
            rightMargin=2 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
            title="Documento GIA",
            author="GestorIA · GIA",
        )
        doc.build(story)

        return GiaService._insert_file(
            cursor,
            conversation_id,
            message_id,
            user_id,
            original_name,
            stored_name,
            "application/pdf",
            path.stat().st_size,
            str(path),
            "pdf",
            content,
        )

    @staticmethod
    def _html_escape(value: str) -> str:
        return (
            value.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    @staticmethod
    def _extract_text(path: Path, mime_type: str) -> str | None:
        
        if mime_type.startswith(TEXT_MIME_PREFIXES) or mime_type in TEXT_MIME_TYPES:
            try:
                return path.read_text(encoding="utf-8", errors="replace")[:20000]
            except Exception as exc:
                logger.warning("No se pudo leer archivo de texto %s: %s", path, exc)
                return None

        
        if mime_type == "application/pdf" or path.suffix.lower() == ".pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
                if not text:
                    return "PDF recibido, pero no contiene texto extraíble (puede ser un PDF escaneado o de solo imagen)."
                return text[:30000]
            except ImportError:
                logger.error("pypdf no está instalado; no se puede leer el PDF")
                return "PDF recibido, pero pypdf no está instalado para extraer su contenido."
            except Exception as exc:
                logger.warning("Error al leer PDF %s: %s", path, exc)
                return "PDF recibido. No se pudo extraer texto automáticamente."

        
        if mime_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
            "application/vnd.oasis.opendocument.text",
        ) or path.suffix.lower() in (".docx", ".doc", ".odt"):
            try:
                import docx
                doc = docx.Document(str(path))
                text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                return text[:25000] if text else "Documento Word recibido, pero no contiene texto extraíble."
            except ImportError:
                logger.warning("python-docx no instalado; no se puede leer el documento Word")
                return "Documento Word recibido, pero python-docx no está instalado para extraer su contenido."
            except Exception as exc:
                logger.warning("Error al leer Word %s: %s", path, exc)
                return "Documento Word recibido. No se pudo extraer el texto automáticamente."

        
        if mime_type in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        ) or path.suffix.lower() in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
                rows: list[str] = []
                for sheet in wb.worksheets:
                    rows.append(f"[Hoja: {sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        line = "\t".join("" if v is None else str(v) for v in row)
                        if line.strip():
                            rows.append(line)
                text = "\n".join(rows)
                return text[:25000] if text else "Hoja Excel recibida, pero no contiene datos extraíbles."
            except ImportError:
                logger.warning("openpyxl no instalado; no se puede leer el Excel")
                return "Hoja Excel recibida, pero openpyxl no está instalado para extraer su contenido."
            except Exception as exc:
                logger.warning("Error al leer Excel %s: %s", path, exc)
                return "Hoja Excel recibida. No se pudo extraer el contenido automáticamente."

        return None

    @staticmethod
    def _attachment_context(attachments: list[dict], conversation_files: list[dict]) -> str:
        
        merged: dict[str, dict] = {}
        for item in (attachments or []):
            if item.get("id"):
                merged[item["id"]] = item
        for item in (conversation_files or []):
            if item.get("id"):
                merged[item["id"]] = item

        if not merged:
            return ""

        files = sorted(
            merged.values(),
            key=lambda x: x.get("created_at") or "",
            reverse=True,
        )

        chunks: list[str] = []
        catalog = []
        for item in files[:25]:
            kind = item.get("tipo", "upload")
            catalog.append(f"- {item.get('nombre_original', 'archivo')} ({kind}, {item.get('mime_type', 'application/octet-stream')})")
        if catalog:
            chunks.append("\n\n[Memoria de archivos de la conversación]\n" + "\n".join(catalog))

        total_chars = 0
        for item in files:
            text = (item.get("extracted_text") or "").strip()
            if text:
                excerpt = text[:5000]
                total_chars += len(excerpt)
                chunks.append(
                    f"\n\n[Archivo (memoria): {item.get('nombre_original', 'archivo')}]\n{excerpt}"
                )
            elif str(item.get("mime_type", "")).startswith("image/"):
                chunks.append(f"\n\n[Imagen en memoria: {item.get('nombre_original', 'imagen')}]")
            else:
                chunks.append(
                    f"\n\n[Archivo en memoria sin texto extraíble: {item.get('nombre_original', 'archivo')}]"
                )
            if total_chars >= 18000:
                chunks.append("\n\n[Nota] Se ha truncado el contexto de archivos por límite de tamaño.")
                break

        return "".join(chunks)

    @staticmethod
    def _image_parts(attachments: list[dict]) -> list[dict]:
        parts = []
        for item in attachments[:6]:
            if not str(item.get("mime_type", "")).startswith(IMAGE_MIME_PREFIXES):
                continue
            if not item.get("ruta_archivo"):
                continue
            path = Path(item["ruta_archivo"])
            if not path.exists():
                continue
            if path.stat().st_size > 8 * 1024 * 1024:
                continue
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
            parts.append({
                "type": "image_url",
                "image_url": {"url": f"data:{item['mime_type']};base64,{encoded}"},
            })
        return parts

    @staticmethod
    def _get_files_for_context(cursor, conversation_id: str) -> list[dict]:
        cursor.execute(
            """
            SELECT
                id::text,
                nombre_original,
                mime_type,
                tamano_bytes,
                tipo,
                ruta_archivo,
                extracted_text,
                created_at
            FROM gia_archivos
            WHERE conversacion_id = %s
            ORDER BY created_at DESC
            LIMIT 25
            """,
            (conversation_id,),
        )
        return [dict(row) for row in cursor.fetchall()]

    @staticmethod
    def _ensure_schema(cursor) -> None:
        cursor.execute('CREATE EXTENSION IF NOT EXISTS "uuid-ossp"')
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS gia_conversaciones (
                id UUID PRIMARY KEY DEFAULT uuid_generate_v4(),
                user_id UUID NOT NULL REFERENCES usuarios(id) ON DELETE CASCADE,
                titulo VARCHAR(180) NOT NULL DEFAULT 'Nueva conversación',
                memoria_resumen TEXT NOT NULL DEFAULT '',
                archivada BOOLEAN NOT NULL DEFAULT FALSE,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """
        )
        cursor.execute(
            """
            ALTER TABLE gia_conversaciones
            ADD COLUMN IF NOT EXISTS memoria_resumen TEXT NOT NULL DEFAULT ''
            """
        )
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS gia_mensajes (
                id UUID PRIMARY KEY DEFAULT uuid_generate_v4(),
                conversacion_id UUID NOT NULL REFERENCES gia_conversaciones(id) ON DELETE CASCADE,
                role VARCHAR(20) NOT NULL,
                content TEXT NOT NULL,
                metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """
        )
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS gia_archivos (
                id UUID PRIMARY KEY DEFAULT uuid_generate_v4(),
                conversacion_id UUID NOT NULL REFERENCES gia_conversaciones(id) ON DELETE CASCADE,
                mensaje_id UUID REFERENCES gia_mensajes(id) ON DELETE SET NULL,
                created_by UUID NOT NULL REFERENCES usuarios(id) ON DELETE CASCADE,
                nombre_original VARCHAR(255) NOT NULL,
                nombre_guardado VARCHAR(255) NOT NULL,
                mime_type VARCHAR(120) NOT NULL,
                tamano_bytes BIGINT NOT NULL DEFAULT 0,
                ruta_archivo TEXT NOT NULL,
                tipo VARCHAR(20) NOT NULL DEFAULT 'upload',
                extracted_text TEXT,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """
        )
        GiaService._ensure_usage_schema(cursor)

    @staticmethod
    def _ensure_usage_schema(cursor) -> None:
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS gia_uso_ia (
                id UUID PRIMARY KEY DEFAULT uuid_generate_v4(),
                user_id UUID REFERENCES usuarios(id) ON DELETE SET NULL,
                nombre_usuario VARCHAR(255) NOT NULL DEFAULT 'desconocido',
                conversacion_id UUID REFERENCES gia_conversaciones(id) ON DELETE SET NULL,
                mensaje_id UUID REFERENCES gia_mensajes(id) ON DELETE SET NULL,
                mode VARCHAR(20) NOT NULL DEFAULT 'respuesta',
                requested_mode VARCHAR(20) NOT NULL DEFAULT 'respuesta',
                model VARCHAR(120),
                prompt_tokens INTEGER NOT NULL DEFAULT 0,
                completion_tokens INTEGER NOT NULL DEFAULT 0,
                total_tokens INTEGER NOT NULL DEFAULT 0,
                image_generations INTEGER NOT NULL DEFAULT 0,
                estimated_cost_eur NUMERIC(12, 6) NOT NULL DEFAULT 0,
                metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """
        )
        cursor.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS idx_gia_uso_ia_mensaje
            ON gia_uso_ia (mensaje_id)
            WHERE mensaje_id IS NOT NULL
            """
        )
        cursor.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_gia_uso_ia_created_at
            ON gia_uso_ia (created_at DESC)
            """
        )
        cursor.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_gia_uso_ia_user_created_at
            ON gia_uso_ia (user_id, created_at DESC)
            """
        )
        cursor.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_gia_uso_ia_mode_created_at
            ON gia_uso_ia (mode, created_at DESC)
            """
        )

    @staticmethod
    def _record_usage(
        cursor,
        user_id: str,
        conversation_id: str,
        message_id: str,
        mode: str,
        requested_mode: str,
        usage: dict,
    ) -> None:
        cursor.execute(
            """
            INSERT INTO gia_uso_ia (
                user_id,
                nombre_usuario,
                conversacion_id,
                mensaje_id,
                mode,
                requested_mode,
                model,
                prompt_tokens,
                completion_tokens,
                total_tokens,
                image_generations,
                estimated_cost_eur,
                metadata,
                created_at
            )
            VALUES (
                %s,
                COALESCE((SELECT usuario::text FROM usuarios WHERE id = %s), 'desconocido'),
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s::jsonb,
                NOW()
            )
            """,
            (
                user_id,
                user_id,
                conversation_id,
                message_id,
                mode,
                requested_mode,
                usage.get("model"),
                int(usage.get("prompt_tokens") or 0),
                int(usage.get("completion_tokens") or 0),
                int(usage.get("total_tokens") or 0),
                int(usage.get("image_generations") or 0),
                float(usage.get("estimated_cost_eur") or 0),
                json.dumps({"source": "gia_send_message"}),
            ),
        )

    @staticmethod
    def _insert_message(cursor, conversation_id: str, role: str, content: str, metadata: dict | None = None) -> dict:
        cursor.execute(
            """
            INSERT INTO gia_mensajes (conversacion_id, role, content, metadata)
            VALUES (%s, %s, %s, %s::jsonb)
            RETURNING id::text, role, content, created_at
            """,
            (conversation_id, role, content, json.dumps(metadata or {})),
        )
        row = dict(cursor.fetchone())
        row["files"] = []
        return row

    @staticmethod
    def _insert_file(
        cursor,
        conversation_id: str,
        message_id: str,
        user_id: str,
        original_name: str,
        stored_name: str,
        mime_type: str,
        size: int,
        path: str,
        kind: str,
        extracted_text: str | None,
    ) -> dict:
        cursor.execute(
            """
            INSERT INTO gia_archivos
                (conversacion_id, mensaje_id, created_by, nombre_original, nombre_guardado,
                 mime_type, tamano_bytes, ruta_archivo, tipo, extracted_text)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            RETURNING id::text, nombre_original, mime_type, tamano_bytes, tipo, created_at
            """,
            (conversation_id, message_id, user_id, original_name, stored_name, mime_type, size, path, kind, extracted_text),
        )
        return GiaService._file_item(cursor.fetchone())

    @staticmethod
    def _get_owned_conversation(cursor, user_id: str, conversation_id: str) -> dict | None:
        cursor.execute(
            """
            SELECT id::text, titulo, memoria_resumen, created_at, updated_at
            FROM gia_conversaciones
            WHERE id = %s AND user_id = %s AND archivada = FALSE
            """,
            (conversation_id, user_id),
        )
        row = cursor.fetchone()
        return dict(row) if row else None

    @staticmethod
    def _get_messages(cursor, conversation_id: str, limit: int | None = None) -> list[dict]:
        if limit:
            cursor.execute(
                """
                SELECT id::text, role, content, created_at
                FROM (
                    SELECT id, role, content, created_at
                    FROM gia_mensajes
                    WHERE conversacion_id = %s
                    ORDER BY created_at DESC
                    LIMIT %s
                ) recent_messages
                ORDER BY created_at ASC
                """,
                (conversation_id, limit),
            )
        else:
            cursor.execute(
                """
                SELECT id::text, role, content, created_at
                FROM gia_mensajes
                WHERE conversacion_id = %s
                ORDER BY created_at ASC
                """,
                (conversation_id,),
            )
        messages = [dict(row) for row in cursor.fetchall()]
        for msg in messages:
            msg["files"] = GiaService._files_for_message(cursor, msg["id"])
        return messages

    @staticmethod
    def _get_files(cursor, conversation_id: str) -> list[dict]:
        cursor.execute(
            """
            SELECT id::text, nombre_original, mime_type, tamano_bytes, tipo, created_at
            FROM gia_archivos
            WHERE conversacion_id = %s
            ORDER BY created_at DESC
            """,
            (conversation_id,),
        )
        return [GiaService._file_item(row) for row in cursor.fetchall()]

    @staticmethod
    def _files_for_message(cursor, message_id: str) -> list[dict]:
        cursor.execute(
            """
            SELECT id::text, nombre_original, mime_type, tamano_bytes, tipo, created_at
            FROM gia_archivos
            WHERE mensaje_id = %s
            ORDER BY created_at ASC
            """,
            (message_id,),
        )
        return [GiaService._file_item(row) for row in cursor.fetchall()]

    @staticmethod
    def _get_conversation_summary(cursor, conversation_id: str) -> dict:
        cursor.execute(
            """
            SELECT
                c.id::text,
                c.titulo,
                c.updated_at,
                COUNT(m.id) AS message_count,
                (
                    SELECT m2.content
                    FROM gia_mensajes m2
                    WHERE m2.conversacion_id = c.id
                    ORDER BY m2.created_at DESC
                    LIMIT 1
                ) AS last_message
            FROM gia_conversaciones c
            LEFT JOIN gia_mensajes m ON m.conversacion_id = c.id
            WHERE c.id = %s
            GROUP BY c.id
            """,
            (conversation_id,),
        )
        return GiaService._conversation_summary(cursor.fetchone())

    @staticmethod
    def _touch_conversation(cursor, conversation_id: str, message: str) -> None:
        title = GiaService._title_from_message(message)
        cursor.execute(
            """
            UPDATE gia_conversaciones
            SET updated_at = NOW(),
                titulo = CASE WHEN titulo = 'Nueva conversación' THEN %s ELSE titulo END
            WHERE id = %s
            """,
            (title, conversation_id),
        )

    @staticmethod
    def _conversation_summary(row) -> dict:
        data = dict(row)
        data["message_count"] = int(data.get("message_count") or 0)
        if data.get("last_message") and len(data["last_message"]) > 120:
            data["last_message"] = data["last_message"][:117] + "..."
        return data

    @staticmethod
    def _should_refresh_memory(message_count: int | None, existing_memory: str | None) -> bool:
        total = int(message_count or 0)
        if total <= 0:
            return False
        if not (existing_memory or "").strip():
            return True
        return total % MEMORY_REFRESH_EVERY_MESSAGES == 0

    @staticmethod
    def _refresh_conversation_memory(cursor, conversation_id: str) -> None:
        cursor.execute(
            """
            SELECT memoria_resumen
            FROM gia_conversaciones
            WHERE id = %s
            """,
            (conversation_id,),
        )
        previous_row = cursor.fetchone()
        previous_memory = (dict(previous_row).get("memoria_resumen") if previous_row else "") or ""

        messages = GiaService._get_messages(cursor, conversation_id, limit=30)
        files = GiaService._get_files_for_context(cursor, conversation_id)
        memory = GiaService._build_conversation_memory(messages, files, previous_memory)
        if len(memory) > MEMORY_LLM_COMPRESSION_TRIGGER:
            memory = GiaService._compress_memory_with_llm(memory)
        cursor.execute(
            """
            UPDATE gia_conversaciones
            SET memoria_resumen = %s
            WHERE id = %s
            """,
            (memory[:MAX_MEMORY_CHARS], conversation_id),
        )

    @staticmethod
    def _build_conversation_memory(messages: list[dict], files: list[dict], previous_memory: str = "") -> str:
        def compact(value: str, limit: int = 220) -> str:
            text = " ".join((value or "").split())
            return text[:limit].rstrip() + ("..." if len(text) > limit else "")

        user_msgs = [compact(m.get("content", "")) for m in messages if m.get("role") == "user" and (m.get("content") or "").strip()]
        assistant_msgs = [compact(m.get("content", "")) for m in messages if m.get("role") == "assistant" and (m.get("content") or "").strip()]

        lines: list[str] = ["Objetivo y estado de la conversación:"]
        if previous_memory.strip():
            lines.append("Memoria previa consolidada:")
            lines.append(f"- {compact(previous_memory, 600)}")

        if user_msgs:
            lines.append(f"- Última petición del usuario: {user_msgs[-1]}")
            if len(user_msgs) > 1:
                lines.append(f"- Petición previa relevante: {user_msgs[-2]}")
        if assistant_msgs:
            lines.append(f"- Última respuesta del asistente: {assistant_msgs[-1]}")

        if files:
            lines.append("Archivos conocidos en esta conversación:")
            for f in files[:10]:
                name = f.get("nombre_original", "archivo")
                mime = f.get("mime_type", "application/octet-stream")
                kind = f.get("tipo", "upload")
                lines.append(f"- {name} ({kind}, {mime})")
                extracted = (f.get("extracted_text") or "").strip()
                if extracted:
                    lines.append(f"  Extracto: {compact(extracted, 180)}")

        if len(user_msgs) > 2:
            lines.append("Contexto adicional reciente:")
            for msg in user_msgs[-5:-1]:
                lines.append(f"- {msg}")

        return "\n".join(lines)[:MAX_MEMORY_CHARS]

    @staticmethod
    def _compress_memory_with_llm(memory: str) -> str:
        if not settings.openai_api_key:
            return memory[:MAX_MEMORY_CHARS]

        try:
            from openai import OpenAI
        except ImportError:
            return memory[:MAX_MEMORY_CHARS]

        try:
            client = OpenAI(api_key=settings.openai_api_key)
            response = client.chat.completions.create(
                model=settings.openai_gia_model,
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "Resume memoria conversacional para uso interno de un asistente. "
                            "Devuelve solo texto en español, claro y factual, sin inventar. "
                            "Estructura recomendada: objetivo, hechos clave, decisiones tomadas, "
                            "archivos relevantes, pendientes. Máximo 1200 caracteres."
                        ),
                    },
                    {
                        "role": "user",
                        "content": memory[:MAX_MEMORY_CHARS],
                    },
                ],
                max_tokens=500,
                temperature=0.1,
            )
            compressed = (response.choices[0].message.content or "").strip()
            if not compressed:
                return memory[:MAX_MEMORY_CHARS]
            return compressed[:MEMORY_LLM_TARGET_CHARS]
        except Exception as exc:
            logger.warning("No se pudo comprimir memoria con LLM: %s", exc)
            return memory[:MAX_MEMORY_CHARS]

    @staticmethod
    def _file_item(row) -> dict:
        data = dict(row)
        data["tamano_bytes"] = int(data.get("tamano_bytes") or 0)
        data["download_url"] = f"/ai/gia/files/{data['id']}/download"
        return data

    @staticmethod
    def _conversation_dir(conversation_id: str) -> Path:
        safe_id = re.sub(r"[^a-zA-Z0-9-]", "", conversation_id)
        path = Path(settings.gia_storage_dir) / safe_id
        path.mkdir(parents=True, exist_ok=True)
        return path

    @staticmethod
    def _title_from_message(message: str) -> str:
        clean = " ".join(message.strip().split())
        return clean[:80] or "Nueva conversación"
